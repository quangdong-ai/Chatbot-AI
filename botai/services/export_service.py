"""Tao file export DOCX/PDF cho cau tra loi va bao cao tai lieu."""

from pathlib import Path

import pymupdf as fitz

__all__ = [
    "_create_docx_export",
    "_create_pdf_export",
    "_create_document_report_docx",
    "_create_document_report_pdf",
]

def _create_docx_export(path: Path, question: str, answer: str, sources: str) -> None:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("Cần cài python-docx để export DOCX.") from exc

    doc = DocxDocument()
    doc.add_heading("Câu trả lời từ chatbot", level=1)
    doc.add_heading("Câu hỏi", level=2)
    doc.add_paragraph(question)
    doc.add_heading("Câu trả lời", level=2)
    for paragraph in answer.split("\n"):
        doc.add_paragraph(paragraph)
    if sources.strip():
        doc.add_heading("Nguồn tham khảo", level=2)
        for paragraph in sources.split("\n"):
            doc.add_paragraph(paragraph)
    doc.save(path)


def _create_pdf_export(path: Path, question: str, answer: str, sources: str) -> None:
    doc = fitz.open()
    page = doc.new_page()
    rect = fitz.Rect(50, 50, 545, 790)
    text = (
        "Câu trả lời từ chatbot\n\n"
        f"Câu hỏi:\n{question}\n\n"
        f"Câu trả lời:\n{answer}\n\n"
        f"Nguồn tham khảo:\n{sources or 'Không có'}"
    )
    page.insert_textbox(rect, text, fontsize=11, fontname="helv", align=0)
    doc.save(path)
    doc.close()


def _create_document_report_docx(path: Path, title: str, body: str, sources: str) -> None:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("Cần cài python-docx để export DOCX.") from exc

    doc = DocxDocument()
    doc.add_heading(title, level=1)
    for paragraph in body.split("\n"):
        text = paragraph.strip()
        if not text:
            continue
        doc.add_paragraph(text)
    if sources.strip():
        doc.add_heading("Nguồn tham khảo", level=2)
        for paragraph in sources.split("\n"):
            doc.add_paragraph(paragraph)
    doc.save(path)


def _create_document_report_pdf(path: Path, title: str, body: str, sources: str) -> None:
    pdf = fitz.open()
    full_text = f"{title}\n\n{body}\n\nNguồn tham khảo:\n{sources or 'Không có'}"
    page = pdf.new_page()
    rect = fitz.Rect(50, 50, 545, 790)
    remaining = full_text
    while remaining:
        written = page.insert_textbox(rect, remaining, fontsize=11, fontname="helv", align=0)
        if written >= 0:
            break
        split_at = max(1200, min(len(remaining), 2600))
        paragraph_break = remaining.rfind("\n", 0, split_at)
        if paragraph_break > 500:
            split_at = paragraph_break
        page.insert_textbox(rect, remaining[:split_at], fontsize=11, fontname="helv", align=0)
        remaining = remaining[split_at:].lstrip()
        if remaining:
            page = pdf.new_page()
    pdf.save(path)
    pdf.close()

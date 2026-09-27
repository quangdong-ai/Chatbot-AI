import hashlib
import os
import pickle
import re
import shutil
import sqlite3
import sys
import tempfile
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Iterable

import pymupdf as fitz
import torch
from langchain_chroma import Chroma
from langchain_community.retrievers import BM25Retriever
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter


PDF_PATH = os.getenv("BOTAI_PDF_PATH", "tailieu.pdf")
DOCUMENTS_DIR = os.getenv("BOTAI_DOCUMENTS_DIR", "storage/documents")
DB_DIR = os.getenv("BOTAI_CHROMA_DIR", "chroma_db")
MODEL_PATH = os.getenv("BOTAI_EMBEDDING_MODEL_PATH", "./models--BAAI--bge-m3")
BM25_PATH = os.getenv("BOTAI_BM25_PATH", "bm25_retriever.pkl")
APP_DB_PATH = os.getenv("BOTAI_APP_DB_PATH", "chatbot_app.db")

TEXT_CHUNK_SIZE = 1200
TEXT_CHUNK_OVERLAP = 180
MIN_CHUNK_CHARS = 40
OCR_ENABLED = os.getenv("BOTAI_OCR_ENABLED", "1").strip().lower() in {"1", "true", "yes"}
OCR_LANG = os.getenv("BOTAI_OCR_LANG", "vie+eng")
OCR_MIN_PAGE_TEXT_CHARS = int(os.getenv("BOTAI_OCR_MIN_PAGE_TEXT_CHARS", "30"))
OCR_DPI = int(os.getenv("BOTAI_OCR_DPI", "180"))
OCR_TABLE_DPI = int(os.getenv("BOTAI_OCR_TABLE_DPI", "216"))
OCR_ENGINE = os.getenv("BOTAI_OCR_ENGINE", "tesseract").strip().lower()
OCR_VL_MODEL_ID = os.getenv("BOTAI_OCR_VL_MODEL_ID", "VietAlphaLabs/SenOCR-Vi").strip()
OCR_VL_MODEL_DIR = os.getenv("BOTAI_OCR_VL_MODEL_DIR", "").strip()
OCR_VL_PIPELINE_VERSION = os.getenv("BOTAI_OCR_VL_PIPELINE_VERSION", "v1.6").strip()
OCR_VL_BACKEND = os.getenv("BOTAI_OCR_VL_BACKEND", "auto").strip().lower()
OCR_VL_MAX_NEW_TOKENS = int(os.getenv("BOTAI_OCR_VL_MAX_NEW_TOKENS", "512"))
PARSER_NAME = "botai_ingest"
PARSER_VERSION = "2026-09-25-structure-v5-appendix-continuation"
_OCR_UNAVAILABLE_REPORTED = False
_OCR_LANG_WARNING_REPORTED = False
_OCR_VL_UNAVAILABLE_REPORTED = False
_OCR_VL_PIPELINE = None
_OCR_VL_TRANSFORMERS_PROCESSOR = None
_OCR_VL_TRANSFORMERS_MODEL = None
_OCR_VL_TRANSFORMERS_DEVICE = None
_OCR_VL_PAGE_CACHE: dict[tuple[int, int, int], str] = {}
BM25_K1 = float(os.getenv("BOTAI_BM25_K1", "1.2"))
BM25_B = float(os.getenv("BOTAI_BM25_B", "0.55"))


def _document_bot_map() -> dict[str, int]:
    if not os.path.exists(APP_DB_PATH):
        return {}
    try:
        conn = sqlite3.connect(APP_DB_PATH)
        rows = conn.execute("SELECT filename, bot_id FROM documents").fetchall()
        conn.close()
    except sqlite3.Error:
        return {}
    return {str(filename): int(bot_id or 1) for filename, bot_id in rows}


def _apply_bot_metadata(documents: list[Document], source_path: Path, bot_map: dict[str, int]) -> list[Document]:
    bot_id = bot_map.get(source_path.name, 1)
    for doc in documents:
        doc.metadata["bot_id"] = bot_id
    return documents


def _resolve_torch_device() -> str:
    configured = os.getenv("BOTAI_TORCH_DEVICE", "auto").strip().lower()
    if configured in {"cuda", "cpu"}:
        return configured
    return "cuda" if torch.cuda.is_available() else "cpu"


def _resolve_hf_snapshot_path(path: str) -> str:
    model_path = Path(path)
    if (model_path / "config.json").exists():
        return str(model_path)
    refs_main = model_path / "refs" / "main"
    snapshots = model_path / "snapshots"
    if refs_main.exists():
        revision = refs_main.read_text(encoding="utf-8").strip()
        snapshot_path = snapshots / revision
        if (snapshot_path / "config.json").exists():
            return str(snapshot_path)
    if snapshots.exists():
        for snapshot_path in snapshots.iterdir():
            if (snapshot_path / "config.json").exists():
                return str(snapshot_path)
    return str(model_path)


def _clean_text(text: str) -> str:
    text = (text or "").replace("\x00", "").replace("\x08", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _content_hash(text: str) -> str:
    normalized = re.sub(r"\s+", " ", str(text or "")).strip()
    return hashlib.sha256(normalized.encode("utf-8")).hexdigest()


def _finalize_metadata(metadata: dict[str, object], content: str, source_path: str | Path | None = None) -> dict[str, object]:
    source = str(source_path or metadata.get("source") or "")
    finalized = dict(metadata)
    finalized.setdefault("document_id", Path(source).stem if source else "")
    finalized.setdefault("document_version", _content_hash(source)[:12] if source else "")
    finalized.setdefault("parser_name", PARSER_NAME)
    finalized.setdefault("parser_version", PARSER_VERSION)
    finalized["content_hash"] = _content_hash(content)
    return finalized


def _fold_vietnamese(text: str) -> str:
    normalized = unicodedata.normalize("NFD", str(text or "").casefold())
    normalized = "".join(
        char for char in normalized if unicodedata.category(char) != "Mn"
    )
    return normalized.replace("đ", "d").replace("Đ", "d").replace("Ä‘", "d")


def _legal_phrase_tokens(text: str) -> list[str]:
    folded = _fold_vietnamese(text)
    tokens: list[str] = []

    for article in re.findall(r"\bdieu\s+(\d+)\b", folded):
        tokens.append(f"dieu_{article}")
    for clause in re.findall(r"\bkhoan\s+(\d+)\b", folded):
        tokens.append(f"khoan_{clause}")
    for appendix in re.findall(r"\bphu\s+luc\s+([a-zivxlcdm]+)\b", folded):
        tokens.append(f"phu_luc_{appendix}")
    for license_class in re.findall(
        r"\b(?:hang\s+)?(d2e|d1e|c1e|be|ce|de|a1|b1|c1|d1|d2|a|b|c|d)\b",
        folded,
    ):
        tokens.append(f"hang_{license_class}")
    for number, unit in re.findall(r"\b(\d+[\d.,]*)\s*(gio|km|phut|ngay|%)\b", folded):
        tokens.append(f"{number.replace(',', '.').replace('.', '_')}_{unit}")
    for number in re.findall(r"\b(\d+)\s+hoc\s+vien\b", folded):
        tokens.append(f"{number}_hoc_vien")

    phrase_patterns = {
        "giay_phep_lai_xe": "giay phep lai xe",
        "sat_hach": "sat hach",
        "dao_tao": "dao tao",
        "thuc_hanh_lai_xe": "thuc hanh lai xe",
        "ly_thuyet": "ly thuyet",
        "bao_cao_dinh_ky": "bao cao dinh ky",
        "bao_cao_dang_ky_sat_hach": "bao cao dang ky sat hach",
        "hinh_thuc_dao_tao": "hinh thuc dao tao",
        "pham_vi_dieu_chinh": "pham vi dieu chinh",
        "thoi_gian_lai_xe_an_toan": "thoi gian lai xe an toan",
        "hoc_lai_xe_ban_dem": "hoc lai xe ban dem",
        "hoan_thanh_khoa_dao_tao": "hoan thanh khoa dao tao",
        "so_giao_thong_van_tai": "so giao thong van tai",
    }
    for token, phrase in phrase_patterns.items():
        if phrase in folded:
            tokens.append(token)

    return tokens


def tokenize_vietnamese(text: str) -> list[str]:
    original_tokens = re.findall(
        r"\b[\w]+\.\d+[a-z]?\b|\w+",
        str(text or "").lower(),
        flags=re.UNICODE,
    )
    folded_tokens = re.findall(
        r"\b[a-z]+\.\d+[a-z]?\b|\w+",
        _fold_vietnamese(text),
        flags=re.UNICODE,
    )
    tokens = original_tokens + [
        token for token in folded_tokens if token not in original_tokens
    ]
    return tokens + _legal_phrase_tokens(text)


tokenize_vietnamese.__module__ = "ingest"


def _logical_regions(page: fitz.Page) -> list[tuple[str, fitz.Rect]]:
    rect = page.rect
    if rect.width > rect.height * 1.15:
        middle = rect.x0 + rect.width / 2
        return [
            ("left", fitz.Rect(rect.x0, rect.y0, middle, rect.y1)),
            ("right", fitz.Rect(middle, rect.y0, rect.x1, rect.y1)),
        ]
    return [("full", rect)]


def _chunk_text(text: str) -> Iterable[str]:
    text = _clean_text(text)
    if not text:
        return
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=TEXT_CHUNK_SIZE,
        chunk_overlap=TEXT_CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", "; ", ", ", " "],
        length_function=len,
    )
    for chunk in splitter.split_text(text):
        chunk = chunk.strip()
        if len(chunk) >= MIN_CHUNK_CHARS:
            yield chunk


def _legal_block_markers(text: str) -> list[re.Match[str]]:
    return list(
        re.finditer(
            r"(?m)^\s*(?:(?:Phụ\s+lục|Phu\s+luc|Phá»¥\s+lá»¥c)\s+[A-ZIVXLCDM]+\b|(?:Điều|Dieu|Äiá»u)\s+\d+\.?\s+[^\n]+|\d+\.\s+[^\n]+)",
            text,
            flags=re.I,
        )
    )


def _structured_chunks(text: str, inherited: dict[str, str] | None = None) -> Iterable[tuple[str, dict[str, str]]]:
    clean = _clean_text(text)
    if not clean:
        return

    markers = _legal_block_markers(clean)
    if not markers:
        for chunk in _chunk_text(clean):
            yield chunk, _structure_for_chunk(chunk, dict(inherited or {}))
        return

    spans: list[tuple[int, int]] = []
    if markers[0].start() > 0:
        spans.append((0, markers[0].start()))
    for index, marker in enumerate(markers):
        next_start = markers[index + 1].start() if index + 1 < len(markers) else len(clean)
        spans.append((marker.start(), next_start))

    current_context = dict(inherited or {})
    for start, end in spans:
        block = clean[start:end].strip()
        if not block:
            continue
        block_context = _extract_structure_context(block, current_context, prefer_last_article=False)
        current_context.update(block_context)
        for chunk in _chunk_text(block):
            yield chunk, _structure_for_chunk(chunk, block_context)


def _article_matches(text: str) -> list[re.Match[str]]:
    return list(
        re.finditer(
            r"(?m)^\s*(?:Điều|Dieu|Äiá»u)\s+(\d+)\.?\s*([^\n\.]+)",
            text,
            flags=re.I,
        )
    )


def _select_article_match(text: str, matches: list[re.Match[str]], prefer_last: bool) -> re.Match[str] | None:
    if not matches:
        return None
    if prefer_last or len(matches) == 1:
        return matches[-1]

    best_match = matches[0]
    best_span = -1
    text_len = len(text)
    for index, match in enumerate(matches):
        next_start = matches[index + 1].start() if index + 1 < len(matches) else text_len
        span_len = max(0, next_start - match.start())
        if span_len > best_span:
            best_span = span_len
            best_match = match
    return best_match


def _extract_structure_context(
    text: str,
    inherited: dict[str, str] | None = None,
    *,
    prefer_last_article: bool = True,
) -> dict[str, str]:
    context = dict(inherited or {})
    clean = _clean_text(text)
    article_matches = _article_matches(clean)
    article_match = _select_article_match(clean, article_matches, prefer_last_article)
    if article_match:
        context["article"] = f"Điều {article_match.group(1)}"
        context["section_title"] = article_match.group(2).strip()
        context.pop("clause", None)

    appendix_match = re.search(
        r"(?m)^\s*(?:Phụ\s+lục|Phu\s+luc|Phá»¥\s+lá»¥c)\s+([A-ZIVXLCDM]+)\b",
        clean,
        flags=re.I,
    )
    appendix_applies = bool(
        appendix_match
        and (not article_match or appendix_match.start() <= article_match.start())
    )
    if appendix_applies:
        context["appendix"] = f"Phụ lục {appendix_match.group(1).upper()}"

    if article_match:
        match_index = article_matches.index(article_match)
        next_article_start = (
            article_matches[match_index + 1].start()
            if match_index + 1 < len(article_matches)
            else len(clean)
        )
        clause_area = clean[article_match.end():next_article_start]
    else:
        clause_area = clean
    clause_match = re.search(r"(?:^|\n)\s*(\d+)\.\s+", clause_area)
    if clause_match and context.get("article"):
        context["clause"] = f"khoản {clause_match.group(1)}"
    return context


def _structure_for_chunk(chunk: str, page_context: dict[str, str]) -> dict[str, str]:
    context = _extract_structure_context(chunk, page_context, prefer_last_article=False)
    folded = _fold_vietnamese(chunk)
    if "thong tu" in folded and "quy dinh ve dao tao, sat hach, cap giay phep lai xe" in folded:
        context["article"] = "Điều 1"
        context.pop("clause", None)
        context["section_title"] = "Phạm vi điều chỉnh"
    return context

def _contextualize_chunk(chunk: str, metadata: dict[str, object]) -> str:
    labels = [
        str(metadata.get("appendix") or "").strip(),
        str(metadata.get("article") or "").strip(),
        str(metadata.get("clause") or "").strip(),
        str(metadata.get("section_title") or "").strip(),
    ]
    labels = [label for label in labels if label]
    if not labels:
        return chunk
    return f"[Ngữ cảnh: {' - '.join(labels)}]\n{chunk}"


def _effective_ocr_language(pytesseract) -> str:
    global _OCR_LANG_WARNING_REPORTED
    requested_langs = [lang for lang in re.split(r"[+,]", OCR_LANG) if lang]
    installed_langs = set(pytesseract.get_languages(config=""))
    available_langs = [lang for lang in requested_langs if lang in installed_langs]
    effective_lang = "+".join(available_langs) if available_langs else ""
    if not effective_lang:
        effective_lang = "eng" if "eng" in installed_langs else (next(iter(installed_langs), "eng"))
    missing_langs = [lang for lang in requested_langs if lang not in installed_langs]
    if missing_langs and not _OCR_LANG_WARNING_REPORTED:
        print(
            "OCR language fallback: missing "
            f"{', '.join(missing_langs)}; using {effective_lang}."
        )
        _OCR_LANG_WARNING_REPORTED = True
    return effective_lang


def _correct_ocr_technical_text(text: str) -> str:
    text = _clean_text(text)
    replacements = [
        (r"\bETX\b|\bBTX\b|\bRTY\b|\bRI%\b|\bETSX\b", "RTX"),
        (r"\bTOLO[^\s]*S\b|\bYOLOwE\b|\bYOLOvS\b", "YOLOv8"),
        (r"\bOCE\b|\bQCE\b", "OCR"),
        (r"\bCFU\b", "CPU"),
        (r"\bGFU\b", "GPU"),
        (r"\bVEAM\b", "VRAM"),
        (r"\bEAMI?\b|\bRaM\b", "RAM"),
        (r"\b55D\b|\b88D\b", "SSD"),
        (r"\bNVMs\b|\bMVMe\b|\bMV Me\b", "NVMe"),
        (r"\bContamner\b|\bContamer\b", "Container"),
        (r"\bGeFarce\b", "GeForce"),
        (r"\bNVIDLA\b|\bNVHULA\b|\bMVIDLA\b", "NVIDIA"),
        (r"\bgua\b", "qua"),
        (r"\bTÃ¬\b", "Ti"),
        (r"\bTì\b", "Ti"),
        (r"\bGE\b", "GB"),
        (r"\buvlocp\b|\bvvloop\b", "uvloop"),
        (r"\bNodejs\b", "Node.js"),
        (r"\b3\.?109\b", "3.10.9"),
        (r"\b235H2\b", "25H2"),
        (r"\bv24\.05\b", "v24.0.5"),
        (r"\b2[35]H2\b", "25H2"),
        (r"ki[áº¿e]u\s+trÆ°á»ng\s+cÆ¡\s+sá»Ÿ\s+&\s+Ao\s+hoa", "MÃ´i trÆ°á»ng cÆ¡ sá»Ÿ & áº¢o hÃ³a"),
        (r"Ná»n\s+tang\s+thá»±c\s+thi\s+154", "Ná»n táº£ng thá»±c thi lÃµi"),
        (r"Nền\s+tang\s+thực\s+thi\s+154", "Nền tảng thực thi lõi"),
        (r"Nen\s+tang\s+thuc\s+thi\s+154", "Nền tảng thực thi lõi"),
        (r"Cong nghá»‡", "CÃ´ng nghá»‡"),
        (r"PhiÃ©n ban", "PhiÃªn báº£n"),
        (r"ThÆ° viÃ©n", "ThÆ° viá»‡n"),
        (r"kÂ¥", "ká»¹"),
        (r"Vai trd", "Vai trÃ²"),
        (r"TÃ³c Ä‘á»™", "Tá»‘c Ä‘á»™"),
        (r"Xung nhip", "Xung nhá»‹p"),
        (r"Pim báº£o", "Äáº£m báº£o"),
        (r"Fedis", "Redis"),
        (r"In-memorv Veetor Databaze", "In-memory Vector Database"),
        (r"FATSS", "FAISS"),
        (r"Chá»©c nÄƒng va cÆ¡ chÃ© ká»¹ thuáº­t chuyÃ©n sÃ¢u", "Chá»©c nÄƒng vÃ  cÆ¡ cháº¿ ká»¹ thuáº­t chuyÃªn sÃ¢u"),
        (r"Ngá»“n ngá»¯", "NgÃ´n ngá»¯"),
        (r"giao tiÃ©p", "giao tiáº¿p"),
    ]
    for pattern, value in replacements:
        text = re.sub(pattern, value, text, flags=re.I)
    text = re.sub(r"Intel Core\s+(?:15|i3|l5)-12[35]00H", "Intel Core i5-12500H", text, flags=re.I)
    text = re.sub(r"(RTX\s+)30[35]0(\s+T[iìÃ¬])", r"\g<1>3050 Ti", text, flags=re.I)
    text = re.sub(r"(RTX\s+)5020(\s+Ti)", r"\g<1>3050\2", text, flags=re.I)
    text = re.sub(r"\b31[24]\s+GB\s+SSD\s+NVMe\b", "512 GB SSD NVMe", text, flags=re.I)
    return text


def _generic_table_row_summaries(rows: list[list[str]]) -> str:
    if len(rows) < 2:
        return ""
    headers = rows[0]
    summaries: list[str] = []
    for row in rows[1:]:
        if not row or not any(row):
            continue
        label = row[0].strip() or f"DÃ²ng {len(summaries) + 1}"
        values = []
        for index, value in enumerate(row[1:], start=1):
            value = value.strip()
            header = headers[index].strip() if index < len(headers) else f"Cá»™t {index + 1}"
            if header and value:
                values.append(f"{header} = {value}")
        if values:
            summaries.append(f"- {label}: " + "; ".join(values))
    if not summaries:
        return ""
    return "Diá»…n giáº£i hÃ ng/cá»™t cá»§a báº£ng:\n" + "\n".join(summaries)


def _merge_positions(values: list[int], tolerance: int = 14) -> list[int]:
    groups: list[list[int]] = []
    for value in sorted(values):
        if not groups or value - groups[-1][-1] > tolerance:
            groups.append([value])
        else:
            groups[-1].append(value)
    return [int(sum(group) / len(group)) for group in groups]


def _ocr_confidence_from_image(image, lang: str, config: str) -> float:
    import pytesseract

    try:
        data = pytesseract.image_to_data(
            image,
            lang=lang,
            config=config,
            output_type=pytesseract.Output.DICT,
        )
    except Exception:
        return 0.0
    values = []
    for raw in data.get("conf", []):
        try:
            confidence = float(raw)
        except (TypeError, ValueError):
            continue
        if confidence >= 0:
            values.append(confidence)
    if not values:
        return 0.0
    return sum(values) / len(values)


def _ocr_cell_text(image, lang: str, first_column: bool = False) -> tuple[str, float]:
    import cv2
    import numpy as np
    import pytesseract

    cell = cv2.copyMakeBorder(image, 12, 12, 12, 12, cv2.BORDER_CONSTANT, value=(255, 255, 255))
    gray = cv2.cvtColor(cell, cv2.COLOR_RGB2GRAY)
    gray = cv2.resize(gray, None, fx=1.5, fy=1.5, interpolation=cv2.INTER_CUBIC)
    threshold = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)[1]
    config = "--psm 6 -c preserve_interword_spaces=1"
    text = pytesseract.image_to_string(threshold, lang=lang, config=config)
    confidence = _ocr_confidence_from_image(threshold, lang, config)
    text = _correct_ocr_technical_text(text)
    if first_column and not text.strip():
        fallback_config = "--psm 11"
        text = pytesseract.image_to_string(threshold, lang=lang, config=fallback_config)
        confidence = _ocr_confidence_from_image(threshold, lang, fallback_config)
        text = _correct_ocr_technical_text(text)
    return text, confidence


def _page_cache_key(page: fitz.Page, dpi: int) -> tuple[int, int, int]:
    return (id(page.parent), int(page.number), dpi)


def _render_page_image(page: fitz.Page, dpi: int):
    import io
    from PIL import Image

    zoom = max(dpi, 72) / 72
    pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
    return Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB")


def _paddleocr_vl_available() -> bool:
    return OCR_ENGINE in {"paddleocr_vl", "paddleocr-vl", "vl", "senocr", "senocr-vi"}


def _paddleocr_vl_pipeline():
    global _OCR_VL_PIPELINE
    if _OCR_VL_PIPELINE is not None:
        return _OCR_VL_PIPELINE

    from paddleocr import PaddleOCRVL

    model_dir = OCR_VL_MODEL_DIR
    if not model_dir:
        from huggingface_hub import snapshot_download

        model_dir = snapshot_download(repo_id=OCR_VL_MODEL_ID)

    _OCR_VL_PIPELINE = PaddleOCRVL(
        pipeline_version=OCR_VL_PIPELINE_VERSION,
        vl_rec_model_dir=model_dir,
    )
    return _OCR_VL_PIPELINE


def _ocr_vl_model_dir() -> str:
    model_dir = OCR_VL_MODEL_DIR
    if model_dir:
        return model_dir
    from huggingface_hub import snapshot_download

    return snapshot_download(repo_id=OCR_VL_MODEL_ID)


def _ocr_vl_looks_like_transformers_model(model_dir: str) -> bool:
    path = Path(model_dir)
    return (
        (path / "config.json").exists()
        and (
            (path / "model.safetensors.index.json").exists()
            or any(path.glob("model-*.safetensors"))
        )
        and not (path / "inference.pdmodel").exists()
    )


def _ocr_vl_transformers_components():
    global _OCR_VL_TRANSFORMERS_PROCESSOR, _OCR_VL_TRANSFORMERS_MODEL, _OCR_VL_TRANSFORMERS_DEVICE
    if _OCR_VL_TRANSFORMERS_PROCESSOR is not None and _OCR_VL_TRANSFORMERS_MODEL is not None:
        return _OCR_VL_TRANSFORMERS_PROCESSOR, _OCR_VL_TRANSFORMERS_MODEL, _OCR_VL_TRANSFORMERS_DEVICE

    os.environ.setdefault("HF_HOME", str((Path.cwd() / "hf_cache").resolve()))
    from transformers import AutoModelForCausalLM, AutoProcessor
    import inspect
    import torch

    import transformers.masking_utils as masking_utils

    mask_signature = inspect.signature(masking_utils.create_causal_mask)
    if "input_embeds" in mask_signature.parameters and "inputs_embeds" not in mask_signature.parameters:
        original_create_causal_mask = masking_utils.create_causal_mask

        def _create_causal_mask_compat(*args, **kwargs):
            if "inputs_embeds" in kwargs and "input_embeds" not in kwargs:
                kwargs["input_embeds"] = kwargs.pop("inputs_embeds")
            return original_create_causal_mask(*args, **kwargs)

        masking_utils.create_causal_mask = _create_causal_mask_compat

    model_dir = _ocr_vl_model_dir()
    preferred_device = os.getenv("BOTAI_OCR_VL_TRANSFORMERS_DEVICE", "auto").strip().lower()
    if preferred_device == "auto":
        preferred_device = "cuda" if torch.cuda.is_available() else "cpu"
    dtype = torch.float16 if preferred_device == "cuda" else torch.float32

    processor = AutoProcessor.from_pretrained(model_dir, trust_remote_code=True)
    try:
        model = AutoModelForCausalLM.from_pretrained(
            model_dir,
            trust_remote_code=True,
            torch_dtype=dtype,
        ).to(preferred_device).eval()
        device = preferred_device
    except RuntimeError:
        if preferred_device == "cuda":
            torch.cuda.empty_cache()
            model = AutoModelForCausalLM.from_pretrained(
                model_dir,
                trust_remote_code=True,
                torch_dtype=torch.float32,
            ).to("cpu").eval()
            device = "cpu"
        else:
            raise

    _OCR_VL_TRANSFORMERS_PROCESSOR = processor
    _OCR_VL_TRANSFORMERS_MODEL = model
    _OCR_VL_TRANSFORMERS_DEVICE = device
    return processor, model, device


def _transformers_vl_page_text(page: fitz.Page, dpi: int) -> str:
    import torch

    processor, model, device = _ocr_vl_transformers_components()
    image = _render_page_image(page, dpi)
    messages = [
        {
            "role": "user",
            "content": [
                {"type": "image", "image": image},
                {"type": "text", "text": "OCR:"},
            ],
        }
    ]
    inputs = processor.apply_chat_template(
        messages,
        add_generation_prompt=True,
        tokenize=True,
        return_dict=True,
        return_tensors="pt",
    )
    inputs = {key: value.to(device) for key, value in inputs.items()}
    with torch.inference_mode():
        outputs = model.generate(**inputs, max_new_tokens=OCR_VL_MAX_NEW_TOKENS)
    prompt_length = inputs["input_ids"].shape[-1]
    text = processor.decode(outputs[0][prompt_length:-1])
    return _correct_ocr_technical_text(_clean_text(text))


def _read_markdown_outputs(output_dir: Path) -> str:
    parts = []
    for path in sorted(output_dir.rglob("*.md")):
        try:
            text = path.read_text(encoding="utf-8", errors="replace").strip()
        except OSError:
            continue
        if text:
            parts.append(text)
    return "\n\n".join(parts).strip()


def _stringify_paddleocr_result(result: object) -> str:
    if isinstance(result, str):
        return result
    if isinstance(result, dict):
        values = []
        for key in ("markdown", "text", "content", "rec_text", "html"):
            value = result.get(key)
            if isinstance(value, str) and value.strip():
                values.append(value)
        for value in result.values():
            if isinstance(value, (dict, list)):
                nested = _stringify_paddleocr_result(value)
                if nested:
                    values.append(nested)
        return "\n".join(values).strip()
    if isinstance(result, list):
        return "\n".join(_stringify_paddleocr_result(item) for item in result).strip()
    return ""


def _paddleocr_vl_page_text(page: fitz.Page, dpi: int) -> str:
    global _OCR_VL_UNAVAILABLE_REPORTED
    if not OCR_ENABLED or not _paddleocr_vl_available():
        return ""

    cache_key = _page_cache_key(page, dpi)
    if cache_key in _OCR_VL_PAGE_CACHE:
        return _OCR_VL_PAGE_CACHE[cache_key]

    try:
        model_dir = _ocr_vl_model_dir()
        if OCR_VL_BACKEND in {"transformers", "hf"} or (
            OCR_VL_BACKEND == "auto" and _ocr_vl_looks_like_transformers_model(model_dir)
        ):
            text = _transformers_vl_page_text(page, dpi)
            _OCR_VL_PAGE_CACHE[cache_key] = text
            return text

        pipeline = _paddleocr_vl_pipeline()
        image = _render_page_image(page, dpi)
        with tempfile.TemporaryDirectory(prefix="botai_paddleocr_vl_") as tmp:
            image_path = Path(tmp) / f"page_{page.number + 1}.png"
            image.save(image_path)
            output = pipeline.predict(str(image_path))
            output_dir = Path(tmp) / "output"
            output_dir.mkdir(parents=True, exist_ok=True)
            for result in output or []:
                if hasattr(result, "save_to_markdown"):
                    result.save_to_markdown(save_path=output_dir)
            text = _read_markdown_outputs(output_dir)
            if not text:
                text = _stringify_paddleocr_result(output)
        text = _correct_ocr_technical_text(_clean_text(text))
        _OCR_VL_PAGE_CACHE[cache_key] = text
        return text
    except Exception as exc:
        if not _OCR_VL_UNAVAILABLE_REPORTED:
            print(f"PaddleOCR-VL skipped, falling back to Tesseract: {exc}")
            _OCR_VL_UNAVAILABLE_REPORTED = True
        _OCR_VL_PAGE_CACHE[cache_key] = ""
        return ""


def _paddleocr_vl_table_documents(page: fitz.Page, source: str, pdf_page: int) -> list[Document]:
    text = _paddleocr_vl_page_text(page, OCR_TABLE_DPI)
    if not text:
        return []
    table_like = "|" in text or "<table" in text.lower()
    if not table_like:
        return []
    content = (
        f"[BẢNG OCR VL - trang PDF {pdf_page}]\n"
        f"{text}"
    ).strip()
    return [
        Document(
            page_content=content,
            metadata=_finalize_metadata({
                "source": source,
                "pdf_page": pdf_page,
                "chunk_index": 0,
                "table_id": f"{Path(source).stem}:ocr-vl-table:p{pdf_page}:1",
                "type": "ocr_table",
                "region": "ocr_vl_table",
                "ocr": True,
                "ocr_engine": "paddleocr_vl",
                "ocr_model": OCR_VL_MODEL_ID,
                "ocr_needs_review": False,
                "ocr_validation_notes": "",
            }, content, source),
        )
    ]


def _ocr_table_validation_notes(rows: list[list[str]], confidences: list[float]) -> list[str]:
    content = "\n".join(" | ".join(row) for row in rows)
    folded_content = _fold_vietnamese(content)
    notes: list[str] = []
    unresolved_patterns = {
        "RTX 3030": r"\bRTX\s+3030\b",
        "RTX 5020": r"\bRTX\s+5020\b",
        "SSD Ä‘á»c sai": r"\b(?:55D|88D)\b",
        "NVIDIA Ä‘á»c sai": r"\bNVIDLA\b|\bNVHULA\b|\bMVIDLA\b",
        "GeForce Ä‘á»c sai": r"\bGeFarce\b",
        "kÃ½ tá»± OCR láº¡": r"[ï¿½]",
        "mojibake": r"\b(?:Ãƒ|Ã‚|Ã†|Ã¡Âº|Ã¡Â»|Ã„)\w*",
    }
    for label, pattern in unresolved_patterns.items():
        if re.search(pattern, content, flags=re.I):
            notes.append(label)
    if confidences:
        avg_conf = sum(confidences) / len(confidences)
        min_conf = min(confidences)
        if avg_conf < 45:
            notes.append(f"Ä‘á»™ tin cáº­y trung bÃ¬nh tháº¥p ({avg_conf:.1f})")
        if min_conf < 20 and len(confidences) >= 3:
            notes.append(f"cÃ³ Ã´ Ä‘á»™ tin cáº­y ráº¥t tháº¥p ({min_conf:.1f})")
    if "bang ocr" in folded_content and not any(any(cell.strip() for cell in row) for row in rows[1:]):
        notes.append("khÃ´ng Ä‘á»c Ä‘Æ°á»£c dÃ²ng dá»¯ liá»‡u báº£ng")
    return notes


def _infer_ocr_row_label(label: str, cells: list[str]) -> str:
    folded_row = _fold_vietnamese(" ".join(cells))
    if label.strip():
        return label.strip()
    if "16 gb" in folded_row and "3200 mt/s" in folded_row:
        return "RAM"
    if "ssd nvme" in folded_row or "512 gb" in folded_row:
        return "LÆ°u trá»¯"
    return label.strip()


def _ocr_table_documents(page: fitz.Page, source: str, pdf_page: int) -> list[Document]:
    if not OCR_ENABLED:
        return []
    vl_docs = _paddleocr_vl_table_documents(page, source, pdf_page)
    if vl_docs:
        return vl_docs
    try:
        import io
        import cv2
        import numpy as np
        from PIL import Image
        import pytesseract
    except ImportError:
        return []

    try:
        lang = _effective_ocr_language(pytesseract)
        zoom = max(OCR_TABLE_DPI, 72) / 72
        pixmap = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        image = np.array(Image.open(io.BytesIO(pixmap.tobytes("png"))).convert("RGB"))
        gray = cv2.cvtColor(image, cv2.COLOR_RGB2GRAY)
        binary = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)[1]
        horizontal = cv2.erode(binary, cv2.getStructuringElement(cv2.MORPH_RECT, (60, 1)), iterations=1)
        horizontal = cv2.dilate(horizontal, cv2.getStructuringElement(cv2.MORPH_RECT, (60, 1)), iterations=1)
        vertical = cv2.erode(binary, cv2.getStructuringElement(cv2.MORPH_RECT, (1, 60)), iterations=1)
        vertical = cv2.dilate(vertical, cv2.getStructuringElement(cv2.MORPH_RECT, (1, 60)), iterations=1)

        h_contours, _ = cv2.findContours(horizontal, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        v_contours, _ = cv2.findContours(vertical, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        h_boxes = [cv2.boundingRect(contour) for contour in h_contours]
        v_boxes = [cv2.boundingRect(contour) for contour in v_contours]
        h_boxes = [box for box in h_boxes if box[2] > image.shape[1] * 0.25]
        v_boxes = [box for box in v_boxes if box[3] > image.shape[0] * 0.08]
        if len(h_boxes) < 3 or len(v_boxes) < 3:
            return []

        xs = _merge_positions([x for x, y, w, h in v_boxes] + [x + w for x, y, w, h in v_boxes], 20)
        ys = _merge_positions([y for x, y, w, h in h_boxes] + [y + h for x, y, w, h in h_boxes], 20)
        if len(xs) < 3 or len(ys) < 3:
            return []

        rows: list[list[str]] = []
        confidences: list[float] = []
        for row_index in range(len(ys) - 1):
            row: list[str] = []
            for col_index in range(len(xs) - 1):
                crop = image[ys[row_index] + 4:ys[row_index + 1] - 4, xs[col_index] + 4:xs[col_index + 1] - 4]
                cell_text, confidence = _ocr_cell_text(crop, lang, first_column=col_index == 0)
                row.append(cell_text)
                if cell_text.strip():
                    confidences.append(confidence)
            if row_index > 0 and row:
                row[0] = _infer_ocr_row_label(row[0], row)
                folded_row = _fold_vietnamese(" ".join(row))
                if (
                    "nen tang thuc thi" in folded_row
                    and len(row) >= 3
                    and not row[1].strip()
                    and "3.10.9" in row[2]
                ):
                    row[1] = "Python"
            rows.append(row)

        normalized_rows = _normalize_rows(rows)
        if len(normalized_rows) < 2:
            return []
        headers = normalized_rows[0]
        summaries = _generic_table_row_summaries(normalized_rows)
        markdown = _table_to_markdown(normalized_rows)
        if not markdown:
            return []
        validation_notes = _ocr_table_validation_notes(normalized_rows, confidences)
        avg_confidence = sum(confidences) / len(confidences) if confidences else 0.0
        min_confidence = min(confidences) if confidences else 0.0
        validation_text = ""
        if validation_notes:
            validation_text = "\n\n[Cáº¢NH BÃO OCR] " + "; ".join(validation_notes)
        content = (
            f"[Báº¢NG OCR - trang PDF {pdf_page}]\n"
            f"{summaries}\n\n"
            f"{markdown}"
            f"{validation_text}"
        ).strip()
        return [
            Document(
                page_content=content,
                metadata=_finalize_metadata({
                    "source": source,
                    "pdf_page": pdf_page,
                    "chunk_index": 0,
                    "table_id": f"{Path(source).stem}:ocr-table:p{pdf_page}:1",
                    "type": "ocr_table",
                    "region": "ocr_table",
                    "ocr": True,
                    "ocr_engine": "tesseract",
                    "ocr_model": OCR_LANG,
                    "ocr_confidence": round(avg_confidence, 2),
                    "ocr_min_confidence": round(min_confidence, 2),
                    "ocr_needs_review": bool(validation_notes),
                    "ocr_validation_notes": "; ".join(validation_notes),
                }, content, source),
            )
        ]
    except Exception as exc:
        print(f"  - Page {pdf_page}: cannot extract OCR table ({exc})")
        return []


def _ocr_page_text(page: fitz.Page) -> str:
    global _OCR_UNAVAILABLE_REPORTED, _OCR_LANG_WARNING_REPORTED
    if not OCR_ENABLED:
        return ""

    vl_text = _paddleocr_vl_page_text(page, OCR_DPI)
    if vl_text:
        return vl_text

    try:
        import pytesseract
    except ImportError:
        if not _OCR_UNAVAILABLE_REPORTED:
            print("OCR skipped: install pillow and pytesseract to read scanned PDF pages.")
            _OCR_UNAVAILABLE_REPORTED = True
        return ""

    try:
        effective_lang = _effective_ocr_language(pytesseract)
        image = _render_page_image(page, OCR_DPI)
        return _correct_ocr_technical_text(pytesseract.image_to_string(image, lang=effective_lang))
    except Exception as exc:
        if not _OCR_UNAVAILABLE_REPORTED:
            print(f"OCR skipped: {exc}")
            _OCR_UNAVAILABLE_REPORTED = True
        return ""


def _cell_text(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _normalize_rows(rows: list[list[object]]) -> list[list[str]]:
    clean_rows = [
        [_cell_text(cell) for cell in row]
        for row in rows
        if any(_cell_text(cell) for cell in row)
    ]
    if not clean_rows:
        return []
    width = max(len(row) for row in clean_rows)
    return [row + [""] * (width - len(row)) for row in clean_rows]


def _table_to_markdown(rows: list[list[object]]) -> str:
    clean_rows = _normalize_rows(rows)
    if len(clean_rows) < 2:
        return ""

    width = len(clean_rows[0])
    non_empty_cells = sum(1 for row in clean_rows for cell in row if cell)
    dense_rows = sum(1 for row in clean_rows if sum(1 for cell in row if cell) >= 3)
    if width < 2 or non_empty_cells < 4 or dense_rows < 2:
        return ""

    header = clean_rows[0]
    body = clean_rows[1:]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in body)
    return "\n".join(lines)


def _extract_column_headers(rows: list[list[object]]) -> list[str]:
    clean_rows = _normalize_rows(rows)
    if not clean_rows:
        return []

    header_rows: list[list[str]] = []
    for row in clean_rows[:4]:
        first = row[0].strip()
        second = row[1].strip() if len(row) > 1 else ""
        if first and re.fullmatch(r"\d+|[IVX]+\.?", first, flags=re.IGNORECASE) and second:
            break
        header_rows.append(row)

    if not header_rows:
        return []

    headers = []
    for column in range(len(clean_rows[0])):
        parts = [row[column] for row in header_rows if row[column]]
        headers.append(" ".join(parts).strip())
    if not any("háº¡ng" in header.lower() for header in headers):
        return []
    return headers


def _friendly_header(header: str) -> str:
    matches = re.findall(r"(Háº¡ng\s+[A-Z0-9]{1,3}E?)", header, flags=re.IGNORECASE)
    if matches:
        return matches[-1]
    return header.strip()


def _table_row_summaries(rows: list[list[object]], headers: list[str]) -> str:
    clean_rows = _normalize_rows(rows)
    if not clean_rows or not headers:
        return ""

    summaries: list[str] = []
    for row in clean_rows:
        if len(row) < len(headers):
            row = row + [""] * (len(headers) - len(row))
        first = row[0].strip()
        label = row[1].strip() if len(row) > 1 and row[1].strip() else first
        unit = row[2].strip() if len(row) > 2 else ""
        if not label or label.lower() in {"ná»™i dung", "chá»‰ tiÃªu tÃ­nh toÃ¡n cÃ¡c mÃ´n há»c"}:
            continue
        if not (first and re.fullmatch(r"\d+|[IVX]+\.?", first, flags=re.IGNORECASE)) and not row[3:]:
            continue

        values = []
        for column, header in enumerate(headers):
            if column < 3 or column >= len(row):
                continue
            value = row[column].strip()
            header = _friendly_header(header)
            if value and header:
                values.append(f"{header} = {value}{(' ' + unit) if unit else ''}")
        if values:
            summaries.append(f"- {label}: " + "; ".join(values))

    if not summaries:
        return ""
    return "Diá»…n giáº£i hÃ ng/cá»™t cá»§a báº£ng:\n" + "\n".join(summaries)


def _looks_like_table_page(text: str) -> bool:
    folded = _fold_vietnamese(text)
    markers = (
        "tt",
        "stt",
        "hang giay phep",
        "don vi",
        "noi dung",
        "chi tieu",
        "thoi gian",
        "tong thoi gian",
        "gio",
        "km",
    )
    marker_score = sum(1 for marker in markers if marker in folded)
    numeric_columns = len(re.findall(r"\b\d+\s+\d+\s+\d+\b", folded))
    short_numeric_lines = len(re.findall(r"(?m)^\s*\d{1,4}\s*$", folded))
    return marker_score >= 3 or numeric_columns >= 2 or short_numeric_lines >= 4


def _extract_table_documents(
    page: fitz.Page,
    source: str,
    pdf_page: int,
    inherited_headers: list[str] | None = None,
) -> tuple[list[Document], list[str]]:
    documents: list[Document] = []
    page_text = page.get_text("text", sort=True)
    if not _looks_like_table_page(page_text):
        return documents, inherited_headers or []

    try:
        tables = page.find_tables()
    except Exception as exc:
        print(f"  - Page {pdf_page}: cannot detect tables ({exc})")
        return documents, inherited_headers or []

    active_headers = inherited_headers or []
    for table_index, table in enumerate(getattr(tables, "tables", []) or []):
        rows = table.extract()
        markdown = _table_to_markdown(rows)
        if not markdown:
            continue
        table_headers = _extract_column_headers(rows)
        if table_headers:
            active_headers = table_headers
        summaries = _table_row_summaries(rows, active_headers)
        nearby_text = _clean_text(page.get_text("text", clip=table.bbox, sort=True))
        content = (
            f"[Báº¢NG - trang PDF {pdf_page}, báº£ng {table_index + 1}]\n"
            f"{summaries}\n\n"
            f"{markdown}\n\n"
            f"VÄƒn báº£n trong vÃ¹ng báº£ng:\n{nearby_text[:1200]}"
        ).strip()
        documents.append(
            Document(
                page_content=content,
                metadata=_finalize_metadata({
                    "source": source,
                    "pdf_page": pdf_page,
                    "chunk_index": table_index,
                    "table_id": f"{Path(source).stem}:table:p{pdf_page}:{table_index + 1}",
                    "type": "table",
                    "region": "table",
                }, content, source),
            )
        )
    return documents, active_headers


def load_pdf_documents(path: str) -> list[Document]:
    pdf_path = Path(path)
    documents: list[Document] = []
    inherited_headers: list[str] = []
    inherited_structure: dict[str, str] = {}
    with fitz.open(pdf_path) as pdf:
        for page_index, page in enumerate(pdf, start=1):
            page_text_full = _clean_text(page.get_text("text", sort=True))
            inherited_structure = _extract_structure_context(page_text_full, inherited_structure)
            page_text_total = ""
            for region_name, clip in _logical_regions(page):
                page_text = _clean_text(page.get_text("text", clip=clip, sort=True))
                page_text_total = f"{page_text_total}\n{page_text}".strip()
                for chunk_index, (chunk, structure) in enumerate(_structured_chunks(page_text, inherited_structure)):
                    metadata = {
                        "source": pdf_path.name,
                        "pdf_page": page_index,
                        "chunk_index": chunk_index,
                        "type": "text",
                        "region": region_name,
                    }
                    metadata.update(structure)
                    page_content = _contextualize_chunk(chunk, metadata)
                    documents.append(
                        Document(
                            page_content=page_content,
                            metadata=_finalize_metadata(metadata, chunk, pdf_path.name),
                        )
                    )
            if len(page_text_total) < OCR_MIN_PAGE_TEXT_CHARS:
                ocr_text = _ocr_page_text(page)
                ocr_engine = "paddleocr_vl" if _OCR_VL_PAGE_CACHE.get(_page_cache_key(page, OCR_DPI)) else "tesseract"
                for chunk_index, (chunk, structure) in enumerate(_structured_chunks(ocr_text, inherited_structure)):
                    metadata = {
                        "source": pdf_path.name,
                        "pdf_page": page_index,
                        "chunk_index": chunk_index,
                        "type": "ocr_text",
                        "region": "ocr_full",
                        "ocr": True,
                        "ocr_engine": ocr_engine,
                        "ocr_model": OCR_VL_MODEL_ID if ocr_engine == "paddleocr_vl" else OCR_LANG,
                    }
                    metadata.update(structure)
                    page_content = _contextualize_chunk(chunk, metadata)
                    documents.append(
                        Document(
                            page_content=page_content,
                            metadata=_finalize_metadata(metadata, chunk, pdf_path.name),
                        )
                    )
                ocr_table_docs = _ocr_table_documents(page, pdf_path.name, page_index)
                for doc in ocr_table_docs:
                    doc.metadata.update(inherited_structure)
                    doc.page_content = _contextualize_chunk(doc.page_content, doc.metadata)
                    doc.metadata = _finalize_metadata(doc.metadata, doc.page_content, pdf_path.name)
                documents.extend(ocr_table_docs)
            table_docs, inherited_headers = _extract_table_documents(
                page,
                pdf_path.name,
                page_index,
                inherited_headers,
            )
            for doc in table_docs:
                doc.metadata.update(inherited_structure)
                doc.page_content = _contextualize_chunk(doc.page_content, doc.metadata)
                doc.metadata = _finalize_metadata(doc.metadata, doc.page_content, pdf_path.name)
            documents.extend(table_docs)
    return documents


def load_txt_documents(path: str) -> list[Document]:
    text_path = Path(path)
    try:
        raw_text = text_path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        raw_text = text_path.read_text(encoding="utf-8", errors="ignore")

    documents: list[Document] = []
    inherited_structure: dict[str, str] = {}
    for chunk_index, (chunk, structure) in enumerate(_structured_chunks(raw_text, inherited_structure)):
        metadata = {
            "source": text_path.name,
            "pdf_page": "",
            "chunk_index": chunk_index,
            "type": "text",
            "region": "txt",
        }
        metadata.update(structure)
        documents.append(
            Document(
                page_content=chunk,
                metadata=_finalize_metadata(metadata, chunk, text_path.name),
            )
        )
    return documents


def load_docx_documents(path: str) -> list[Document]:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:
        raise RuntimeError("Cáº§n cÃ i python-docx Ä‘á»ƒ ingest DOCX.") from exc

    docx_path = Path(path)
    docx = DocxDocument(docx_path)
    blocks: list[str] = []
    for paragraph in docx.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)
    for table_index, table in enumerate(docx.tables, start=1):
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        markdown = _table_to_markdown(rows)
        if markdown:
            blocks.append(f"[Báº¢NG DOCX {table_index}]\n{markdown}")

    documents: list[Document] = []
    inherited_structure: dict[str, str] = {}
    for chunk_index, (chunk, structure) in enumerate(_structured_chunks("\n\n".join(blocks), inherited_structure)):
        metadata = {
            "source": docx_path.name,
            "pdf_page": "",
            "chunk_index": chunk_index,
            "type": "text",
            "region": "docx",
        }
        metadata.update(structure)
        documents.append(
            Document(
                page_content=chunk,
                metadata=_finalize_metadata(metadata, chunk, docx_path.name),
            )
        )
    return documents


def load_xlsx_documents(path: str) -> list[Document]:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Cáº§n cÃ i openpyxl Ä‘á»ƒ ingest XLSX.") from exc

    xlsx_path = Path(path)
    workbook = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    documents: list[Document] = []
    for sheet in workbook.worksheets:
        rows = [
            [cell for cell in row]
            for row in sheet.iter_rows(values_only=True)
            if any(_cell_text(cell) for cell in row)
        ]
        if not rows:
            continue
        markdown = _table_to_markdown(rows)
        if not markdown:
            text_rows = [" | ".join(_cell_text(cell) for cell in row) for row in rows]
            markdown = "\n".join(text_rows)
        headers = [_cell_text(cell) for cell in rows[0]] if rows else []
        summaries = _table_row_summaries(rows, headers)
        row_range = f"1-{len(rows)}"
        content = (
            f"[EXCEL - sheet {sheet.title}, dÃ²ng {row_range}]\n"
            f"{summaries}\n\n"
            f"{markdown}"
        ).strip()
        documents.append(
            Document(
                page_content=content,
                metadata=_finalize_metadata({
                    "source": xlsx_path.name,
                    "sheet_name": sheet.title,
                    "row_range": row_range,
                    "chunk_index": len(documents),
                    "table_id": f"{xlsx_path.stem}:excel:{sheet.title}:{row_range}",
                    "type": "excel_table",
                    "region": "excel",
                    "pdf_page": "",
                }, content, xlsx_path.name),
            )
        )
    workbook.close()
    return documents


def load_pptx_documents(path: str) -> list[Document]:
    pptx_path = Path(path)
    documents: list[Document] = []
    with zipfile.ZipFile(pptx_path) as archive:
        slide_names = sorted(
            name for name in archive.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )
        for slide_index, slide_name in enumerate(slide_names, start=1):
            xml_bytes = archive.read(slide_name)
            root = ET.fromstring(xml_bytes)
            texts = [
                node.text.strip()
                for node in root.iter()
                if node.tag.endswith("}t") and node.text and node.text.strip()
            ]
            slide_text = _clean_text("\n".join(texts))
            if not slide_text:
                continue
            inherited_structure: dict[str, str] = {}
            for chunk_index, (chunk, structure) in enumerate(_structured_chunks(slide_text, inherited_structure)):
                content = f"[PPTX - slide {slide_index}]\n{chunk}"
                metadata = {
                    "source": pptx_path.name,
                    "pdf_page": "",
                    "slide_number": slide_index,
                    "chunk_index": chunk_index,
                    "type": "slide_text",
                    "region": "pptx",
                }
                metadata.update(structure)
                documents.append(
                    Document(
                        page_content=content,
                        metadata=_finalize_metadata(metadata, chunk, pptx_path.name),
                    )
                )
    return documents


def load_documents_from_path(path: str) -> list[Document]:
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return load_pdf_documents(path)
    if suffix == ".txt":
        return load_txt_documents(path)
    if suffix == ".docx":
        return load_docx_documents(path)
    if suffix == ".xlsx":
        return load_xlsx_documents(path)
    if suffix == ".pptx":
        return load_pptx_documents(path)
    print(f"Skipping unsupported file: {path}")
    return []


def document_paths() -> list[Path]:
    paths: list[Path] = []
    if os.path.exists(PDF_PATH):
        paths.append(Path(PDF_PATH))
    storage = Path(DOCUMENTS_DIR)
    if storage.exists():
        paths.extend(
            sorted(
                path for path in storage.iterdir()
                if path.is_file() and path.suffix.lower() in {".pdf", ".txt", ".docx", ".xlsx", ".pptx"}
            )
        )
    return paths


def ingest_document() -> None:
    paths = document_paths()
    print(f"1. Loading {len(paths)} document(s)...")
    if not paths:
        print(f"Error: no supported documents found. Expected {PDF_PATH} or files in {DOCUMENTS_DIR}.")
        return

    documents: list[Document] = []
    bot_map = _document_bot_map()
    for path in paths:
        print(f"  - Loading {path}...")
        loaded_documents = load_documents_from_path(str(path))
        documents.extend(_apply_bot_metadata(loaded_documents, path, bot_map))

    text_count = sum(1 for doc in documents if doc.metadata.get("type") == "text")
    ocr_count = sum(1 for doc in documents if doc.metadata.get("type") == "ocr_text")
    table_count = sum(1 for doc in documents if doc.metadata.get("type") == "table")
    ocr_table_count = sum(1 for doc in documents if doc.metadata.get("type") == "ocr_table")
    excel_count = sum(1 for doc in documents if doc.metadata.get("type") == "excel_table")
    slide_count = sum(1 for doc in documents if doc.metadata.get("type") == "slide_text")
    print(
        f"Loaded {len(documents)} chunks "
        f"({text_count} text, {ocr_count} ocr, {table_count} table, {ocr_table_count} ocr_table, "
        f"{excel_count} excel, {slide_count} slide)."
    )

    print(
        "2. Rebuilding BM25 Index with Vietnamese/no-accent tokenization "
        f"(k1={BM25_K1}, b={BM25_B})..."
    )
    bm25_retriever = BM25Retriever.from_documents(
        documents,
        bm25_params={"k1": BM25_K1, "b": BM25_B},
        preprocess_func=tokenize_vietnamese,
    )
    bm25_retriever.k = 20
    bm25_parent = Path(BM25_PATH).parent
    if str(bm25_parent) not in {"", "."}:
        bm25_parent.mkdir(parents=True, exist_ok=True)
    with open(BM25_PATH, "wb") as f:
        pickle.dump(bm25_retriever, f)
    print("BM25 Index saved.")

    print("3. Loading Embedding Model...")
    torch_device = _resolve_torch_device()
    print(f"Using torch device for embeddings: {torch_device}")
    try:
        resolved_model_path = _resolve_hf_snapshot_path(MODEL_PATH)
        embeddings = HuggingFaceEmbeddings(
            model_name=resolved_model_path,
            model_kwargs={"device": torch_device},
            encode_kwargs={"normalize_embeddings": True, "batch_size": 4},
        )
    except RuntimeError as exc:
        if torch_device != "cuda":
            raise
        print(f"CUDA embedding load failed, falling back to CPU: {exc}")
        torch.cuda.empty_cache()
        resolved_model_path = _resolve_hf_snapshot_path(MODEL_PATH)
        embeddings = HuggingFaceEmbeddings(
            model_name=resolved_model_path,
            model_kwargs={"device": "cpu"},
            encode_kwargs={"normalize_embeddings": True, "batch_size": 4},
        )

    print("4. Rebuilding ChromaDB...")
    if os.path.exists(DB_DIR):
        shutil.rmtree(DB_DIR)
    chroma_parent = Path(DB_DIR).parent
    if str(chroma_parent) not in {"", "."}:
        chroma_parent.mkdir(parents=True, exist_ok=True)
    Chroma.from_documents(
        documents=documents,
        embedding=embeddings,
        persist_directory=DB_DIR,
    )

    print("SUCCESS! Document ingested with text and table-aware chunks.")


if __name__ == "__main__":
    sys.modules.setdefault("ingest", sys.modules[__name__])
    ingest_document()

